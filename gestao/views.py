from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.utils import timezone
from django.db.models import Sum, Count, Q
from django.http import JsonResponse, HttpResponse
import json
from decimal import Decimal
from datetime import date, timedelta
import traceback
from .models import Cliente, Veiculo, Locacao, Pagamento, Manutencao, Despesa, LogAcesso, Usuario, Peca
from .forms import (
    ClienteForm, VeiculoForm, LocacaoForm, PagamentoForm, PagamentoEditForm,
    ManutencaoForm, DespesaForm, LocacaoFinalizarForm, PecaForm,
)


def get_alertas():
    hoje = date.today()
    alertas = []

    # Pagamentos do dia
    pag_hoje = Pagamento.objects.filter(data_vencimento=hoje, situacao='pendente')
    for p in pag_hoje:
        alertas.append({'tipo': 'warning', 'mensagem': f'Pagamento de R$ {p.valor} vence hoje - {p.locacao.cliente.nome}', 'categoria': 'pagamento'})

    # Pagamentos atrasados
    pag_atrasados = Pagamento.objects.filter(data_vencimento__lt=hoje, situacao__in=['pendente', 'atrasado'])
    for p in pag_atrasados:
        p.situacao = 'atrasado'
        p.save()
        alertas.append({'tipo': 'danger', 'mensagem': f'Pagamento ATRASADO: R$ {p.valor} - {p.locacao.cliente.nome}', 'categoria': 'pagamento'})

    # Manutenções próximas (30 dias)
    manutencoes = Manutencao.objects.all().order_by('-data')
    revisoes_check = {}
    for m in manutencoes:
        vid = m.veiculo_id
        if vid not in revisoes_check:
            revisoes_check[vid] = m
            if m.proxima_revisao and m.proxima_revisao <= hoje + timedelta(days=30):
                dias = (m.proxima_revisao - hoje).days
                txt = f'vence em {dias} dias' if dias >= 0 else 'VENCIDA'
                tipo = 'warning' if dias >= 0 else 'danger'
                alertas.append({'tipo': tipo, 'mensagem': f'Revisão de {m.veiculo} {txt}', 'categoria': 'manutencao'})
            if m.vencimento_seguro and m.vencimento_seguro <= hoje + timedelta(days=30):
                dias = (m.vencimento_seguro - hoje).days
                txt = f'vence em {dias} dias' if dias >= 0 else 'VENCIDO'
                tipo = 'warning' if dias >= 0 else 'danger'
                alertas.append({'tipo': tipo, 'mensagem': f'Seguro de {m.veiculo} {txt}', 'categoria': 'seguro'})
            if m.vencimento_ipva and m.vencimento_ipva <= hoje + timedelta(days=30):
                dias = (m.vencimento_ipva - hoje).days
                txt = f'vence em {dias} dias' if dias >= 0 else 'VENCIDO'
                tipo = 'warning' if dias >= 0 else 'danger'
                alertas.append({'tipo': tipo, 'mensagem': f'IPVA de {m.veiculo} {txt}', 'categoria': 'ipva'})
            if m.vencimento_licenciamento and m.vencimento_licenciamento <= hoje + timedelta(days=30):
                dias = (m.vencimento_licenciamento - hoje).days
                txt = f'vence em {dias} dias' if dias >= 0 else 'VENCIDO'
                tipo = 'warning' if dias >= 0 else 'danger'
                alertas.append({'tipo': tipo, 'mensagem': f'Licenciamento de {m.veiculo} {txt}', 'categoria': 'licenciamento'})

    return alertas


def login_view(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        if user:
            login(request, user)
            LogAcesso.objects.create(usuario=user, acao='login', ip=request.META.get('REMOTE_ADDR'))
            return redirect('dashboard')
        else:
            messages.error(request, 'Usuário ou senha inválidos.')
    return render(request, 'gestao/login.html')


@login_required
def dashboard(request):
    try:
        hoje = date.today()

        mes_param = request.GET.get('mes')
        if mes_param:
            try:
                ano_atual, mes_atual = map(int, mes_param.split('-'))
            except (ValueError, AttributeError):
                mes_atual, ano_atual = hoje.month, hoje.year
        else:
            mes_atual, ano_atual = hoje.month, hoje.year

        mes_selecionado = f"{ano_atual:04d}-{mes_atual:02d}"

        receita_mes = Pagamento.objects.filter(
            data_pagamento__month=mes_atual,
            data_pagamento__year=ano_atual,
            situacao='pago'
        ).aggregate(total=Sum('valor'))['total'] or Decimal('0')

        gastos_manutencao_mes = Manutencao.objects.filter(
            data__month=mes_atual,
            data__year=ano_atual
        ).aggregate(total=Sum('valor'))['total'] or Decimal('0')

        gastos_despesa_mes = Despesa.objects.filter(
            data__month=mes_atual,
            data__year=ano_atual
        ).aggregate(total=Sum('valor'))['total'] or Decimal('0')

        gastos_mes = gastos_manutencao_mes + gastos_despesa_mes
        lucro_mes = receita_mes - gastos_mes

        pag_hoje = Pagamento.objects.filter(
            data_vencimento=hoje,
            situacao='pendente'
        ).count()

        pag_atrasados = Pagamento.objects.filter(
            data_vencimento__lt=hoje,
            situacao__in=['pendente', 'atrasado']
        ).count()

        veiculos_alugados = Veiculo.objects.filter(status='alugado').count()
        veiculos_disponiveis = Veiculo.objects.filter(status='disponivel').count()
        veiculos_manutencao = Veiculo.objects.filter(status='manutencao').count()

        meses_labels = []
        receitas_data = []
        gastos_data = []

        for i in range(5, -1, -1):
            d = hoje - timedelta(days=30 * i)
            meses_labels.append(d.strftime('%b/%y'))

            r = Pagamento.objects.filter(
                data_pagamento__month=d.month,
                data_pagamento__year=d.year,
                situacao='pago'
            ).aggregate(total=Sum('valor'))['total'] or Decimal('0')

            g_manutencao = Manutencao.objects.filter(
                data__month=d.month, data__year=d.year
            ).aggregate(total=Sum('valor'))['total'] or Decimal('0')

            g_despesa = Despesa.objects.filter(
                data__month=d.month, data__year=d.year
            ).aggregate(total=Sum('valor'))['total'] or Decimal('0')

            receitas_data.append(float(r))
            gastos_data.append(float(g_manutencao + g_despesa))

        try:
            locacoes_recentes = (
                Locacao.objects
                .select_related('cliente', 'veiculo')
                .filter(status='ativa')[:5]
            )
        except Exception:
            traceback.print_exc()
            locacoes_recentes = []

        try:
            alertas = get_alertas()
        except Exception:
            alertas = []

        ctx = {
            'receita_mes': receita_mes,
            'gastos_mes': gastos_mes,
            'lucro_mes': lucro_mes,
            'pag_hoje': pag_hoje,
            'pag_atrasados': pag_atrasados,
            'veiculos_alugados': veiculos_alugados,
            'veiculos_disponiveis': veiculos_disponiveis,
            'veiculos_manutencao': veiculos_manutencao,
            'locacoes_recentes': locacoes_recentes,
            'alertas': alertas,
            'meses_labels': json.dumps(meses_labels),
            'receitas_data': json.dumps(receitas_data),
            'gastos_data': json.dumps(gastos_data),
            'mes_selecionado': mes_selecionado,
        }

        return render(request, 'gestao/dashboard.html', ctx)

    except Exception:
        traceback.print_exc()
        raise


@login_required
def clientes_lista(request):
    q = request.GET.get('q', '')
    clientes = Cliente.objects.filter(ativo=True)
    if q:
        clientes = clientes.filter(Q(nome__icontains=q) | Q(sobrenome__icontains=q) | Q(cpf__icontains=q))
    return render(request, 'gestao/clientes_lista.html', {'clientes': clientes, 'q': q})


@login_required
def cliente_form(request, pk=None):
    cliente = get_object_or_404(Cliente, pk=pk) if pk else None
    form = ClienteForm(request.POST or None, request.FILES or None, instance=cliente)
    if form.is_valid():
        form.save()
        messages.success(request, 'Cliente salvo com sucesso!')
        return redirect('clientes_lista')
    return render(request, 'gestao/cliente_form.html', {'form': form, 'cliente': cliente})


@login_required
def cliente_detalhe(request, pk):
    cliente = get_object_or_404(Cliente, pk=pk)
    locacoes = cliente.locacoes.select_related('veiculo').order_by('-data_retirada')
    return render(request, 'gestao/cliente_detalhe.html', {'cliente': cliente, 'locacoes': locacoes})


@login_required
def cliente_excluir(request, pk):
    cliente = get_object_or_404(Cliente, pk=pk)
    if request.method == 'POST':
        cliente.ativo = False
        cliente.save()
        messages.success(request, 'Cliente removido.')
        return redirect('clientes_lista')
    return render(request, 'gestao/confirmar_exclusao.html', {'objeto': cliente, 'tipo': 'cliente'})


@login_required
def veiculos_lista(request):
    status = request.GET.get('status', '')
    veiculos = Veiculo.objects.filter(ativo=True)
    if status:
        veiculos = veiculos.filter(status=status)
    return render(request, 'gestao/veiculos_lista.html', {'veiculos': veiculos, 'status': status})


@login_required
def veiculo_form(request, pk=None):
    veiculo = get_object_or_404(Veiculo, pk=pk) if pk else None
    form = VeiculoForm(request.POST or None, request.FILES or None, instance=veiculo)
    if form.is_valid():
        form.save()
        messages.success(request, 'Veículo salvo com sucesso!')
        return redirect('veiculos_lista')
    return render(request, 'gestao/veiculo_form.html', {'form': form, 'veiculo': veiculo})


@login_required
def veiculo_detalhe(request, pk):
    veiculo = get_object_or_404(Veiculo, pk=pk)
    locacoes = veiculo.locacoes.select_related('cliente').order_by('-data_retirada')
    manutencoes = veiculo.manutencoes.order_by('-data')
    despesas = veiculo.despesas.order_by('-data')

    total_manutencao = manutencoes.aggregate(total=Sum('valor'))['total'] or Decimal('0')
    total_despesa = despesas.aggregate(total=Sum('valor'))['total'] or Decimal('0')
    total_gasto = total_manutencao + total_despesa

    return render(request, 'gestao/veiculo_detalhe.html', {
        'veiculo': veiculo,
        'locacoes': locacoes,
        'manutencoes': manutencoes,
        'despesas': despesas,
        'total_manutencao': total_manutencao,
        'total_despesa': total_despesa,
        'total_gasto': total_gasto,
    })


@login_required
def veiculo_excluir(request, pk):
    veiculo = get_object_or_404(Veiculo, pk=pk)
    if request.method == 'POST':
        veiculo.ativo = False
        veiculo.save()
        messages.success(request, 'Veículo removido.')
        return redirect('veiculos_lista')
    return render(request, 'gestao/confirmar_exclusao.html', {'objeto': veiculo, 'tipo': 'veiculo'})


@login_required
def locacao_renovar(request, pk):
    locacao = get_object_or_404(Locacao, pk=pk)
    if request.method == 'POST':
        hoje = date.today()
        pagamento_pendente = locacao.pagamentos.filter(
            situacao__in=['pendente', 'atrasado']
        ).order_by('data_vencimento').first()

        if pagamento_pendente:
            pagamento_pendente.situacao = 'pago'
            pagamento_pendente.data_pagamento = hoje
            pagamento_pendente.save()
            base_data = pagamento_pendente.data_vencimento
        else:
            base_data = hoje

        proximo_vencimento = base_data + timedelta(days=locacao.periodicidade_dias)
        Pagamento.objects.create(
            locacao=locacao,
            valor=locacao.valor_combinado,
            data_vencimento=proximo_vencimento,
            forma_pagamento=locacao.forma_pagamento,
            situacao='pendente'
        )
        messages.success(request, f'Renovado! Próximo pagamento em {proximo_vencimento.strftime("%d/%m/%Y")}.')

    return redirect('locacoes_lista')


@login_required
def locacoes_lista(request):
    status = request.GET.get('status', '')
    locacoes = Locacao.objects.select_related('cliente', 'veiculo').all()
    if status:
        locacoes = locacoes.filter(status=status)
    return render(request, 'gestao/locacoes_lista.html', {'locacoes': locacoes, 'status': status})


@login_required
def locacao_form(request, pk=None):
    locacao = get_object_or_404(Locacao, pk=pk) if pk else None
    form = LocacaoForm(request.POST or None, instance=locacao)
    if form.is_valid():
        obj = form.save(commit=False)
        obj.criado_por = request.user
        obj.save()

        # Atualiza status do veículo
        obj.veiculo.status = 'alugado'
        obj.veiculo.save()

        # Cria primeiro ciclo de pagamento (a cada X dias)
        if not pk:
            Pagamento.objects.create(
                locacao=obj,
                valor=obj.valor_combinado,
                data_vencimento=obj.data_retirada,
                data_pagamento=obj.data_retirada,
                forma_pagamento=obj.forma_pagamento,
                situacao='pago'
            )
            Pagamento.objects.create(
                locacao=obj,
                valor=obj.valor_combinado,
                data_vencimento=obj.data_retirada + timedelta(days=obj.periodicidade_dias),
                forma_pagamento=obj.forma_pagamento,
                situacao='pendente'
            )
        messages.success(request, 'Locação registrada com sucesso!')
        return redirect('locacoes_lista')
    return render(request, 'gestao/locacao_form.html', {'form': form, 'locacao': locacao})


@login_required
def locacao_detalhe(request, pk):
    locacao = get_object_or_404(Locacao, pk=pk)
    pagamentos = locacao.pagamentos.all()
    return render(request, 'gestao/locacao_detalhe.html', {'locacao': locacao, 'pagamentos': pagamentos})


@login_required
def locacao_finalizar(request, pk):
    locacao = get_object_or_404(Locacao, pk=pk)
    form = LocacaoFinalizarForm(request.POST or None, instance=locacao)
    if form.is_valid():
        obj = form.save(commit=False)
        obj.status = 'finalizada'
        obj.save()
        obj.veiculo.status = 'disponivel'
        obj.veiculo.save()
        messages.success(request, 'Locação finalizada!')
        return redirect('locacoes_lista')
    return render(request, 'gestao/locacao_finalizar.html', {'form': form, 'locacao': locacao})


@login_required
def locacao_excluir(request, pk):
    locacao = get_object_or_404(Locacao, pk=pk)
    if request.method == 'POST':
        if locacao.status == 'ativa':
            locacao.veiculo.status = 'disponivel'
            locacao.veiculo.save()
        locacao.delete()
        messages.success(request, 'Locação excluída (junto com seus pagamentos).')
    return redirect('locacoes_lista')


@login_required
def locacao_editar_inicio(request, pk):
    locacao = get_object_or_404(Locacao, pk=pk)
    if request.method == 'POST':
        data_str = request.POST.get('data_inicio_real')
        locacao.data_inicio_real = data_str or None
        locacao.save()
        messages.success(request, 'Data de início atualizada!')
    return redirect('locacoes_lista')


@login_required
def pagamentos_lista(request):
    situacao = request.GET.get('situacao', '')
    pagamentos = Pagamento.objects.select_related('locacao__cliente', 'locacao__veiculo').all()
    if situacao:
        pagamentos = pagamentos.filter(situacao=situacao)
    # Atualizar atrasados
    hoje = date.today()
    pagamentos.filter(data_vencimento__lt=hoje, situacao='pendente').update(situacao='atrasado')
    return render(request, 'gestao/pagamentos_lista.html', {'pagamentos': pagamentos, 'situacao': situacao})


@login_required
def pagamento_form(request):
    form = PagamentoForm(request.POST or None)
    if form.is_valid():
        form.save()
        messages.success(request, 'Pagamento registrado!')
        return redirect('pagamentos_lista')
    return render(request, 'gestao/pagamento_form.html', {'form': form})


@login_required
def pagamento_editar(request, pk):
    pagamento = get_object_or_404(Pagamento, pk=pk)
    form = PagamentoEditForm(request.POST or None, instance=pagamento)
    if form.is_valid():
        form.save()
        messages.success(request, 'Pagamento atualizado com sucesso!')
        return redirect('pagamentos_lista')
    return render(request, 'gestao/pagamento_editar.html', {'form': form, 'pagamento': pagamento})


@login_required
def pagamento_pagar(request, pk):
    pagamento = get_object_or_404(Pagamento, pk=pk)
    if request.method == 'POST':
        pagamento.situacao = 'pago'
        pagamento.data_pagamento = date.today()
        pagamento.save()

        locacao = pagamento.locacao
        if locacao.status == 'ativa':
            proximo_vencimento = pagamento.data_vencimento + timedelta(days=locacao.periodicidade_dias)
            Pagamento.objects.create(
                locacao=locacao,
                valor=locacao.valor_combinado,
                data_vencimento=proximo_vencimento,
                forma_pagamento=locacao.forma_pagamento,
                situacao='pendente'
            )
            messages.success(request, f'Pagamento confirmado! Próximo ciclo vence em {proximo_vencimento.strftime("%d/%m/%Y")}.')
        else:
            messages.success(request, 'Pagamento marcado como pago!')

    return redirect('pagamentos_lista')


@login_required
def pagamento_excluir(request, pk):
    pagamento = get_object_or_404(Pagamento, pk=pk)
    if request.method == 'POST':
        pagamento.delete()
        messages.success(request, 'Pagamento removido.')
    return redirect('pagamentos_lista')


@login_required
def manutencoes_lista(request):
    veiculo_id = request.GET.get('veiculo', '')
    manutencoes = Manutencao.objects.select_related('veiculo').all()
    if veiculo_id:
        manutencoes = manutencoes.filter(veiculo_id=veiculo_id)
    veiculos = Veiculo.objects.filter(ativo=True)
    return render(request, 'gestao/manutencoes_lista.html', {'manutencoes': manutencoes, 'veiculos': veiculos, 'veiculo_id': veiculo_id})


@login_required
def manutencao_form(request, pk=None):
    manutencao = get_object_or_404(Manutencao, pk=pk) if pk else None
    form = ManutencaoForm(request.POST or None, instance=manutencao)
    if form.is_valid():
        form.save()
        messages.success(request, 'Manutenção registrada!')
        return redirect('manutencoes_lista')
    return render(request, 'gestao/manutencao_form.html', {'form': form, 'manutencao': manutencao})


@login_required
def despesas_lista(request):
    veiculo_id = request.GET.get('veiculo', '')

    despesas = Despesa.objects.select_related('veiculo').order_by('-data')

    if veiculo_id:
        despesas = despesas.filter(veiculo_id=veiculo_id)

    total = despesas.aggregate(total=Sum('valor'))['total'] or Decimal('0')

    veiculos = Veiculo.objects.filter(ativo=True)

    return render(request, 'gestao/despesas_lista.html', {
        'despesas': despesas,
        'veiculos': veiculos,
        'veiculo_id': veiculo_id,
        'total': total,
    })


@login_required
def despesa_form(request, pk=None):
    despesa = get_object_or_404(Despesa, pk=pk) if pk else None
    form = DespesaForm(request.POST or None, instance=despesa)

    if form.is_valid():
        form.save()
        messages.success(request, 'Despesa salva com sucesso!')
        return redirect('despesas_lista')

    return render(request, 'gestao/despesa_form.html', {'form': form, 'despesa': despesa})


@login_required
def despesa_excluir(request, pk):
    despesa = get_object_or_404(Despesa, pk=pk)

    if request.method == 'POST':
        despesa.delete()
        messages.success(request, 'Despesa excluída com sucesso.')
        return redirect('despesas_lista')

    return render(request, 'gestao/despesa_excluir.html', {'despesa': despesa})


@login_required
def financeiro(request):
    hoje = date.today()
    mes = int(request.GET.get('mes', hoje.month))
    ano = int(request.GET.get('ano', hoje.year))

    receitas = Pagamento.objects.filter(
        data_pagamento__month=mes,
        data_pagamento__year=ano,
        situacao='pago'
    ).select_related('locacao__cliente')

    gastos_manutencao = Manutencao.objects.filter(
        data__month=mes,
        data__year=ano
    ).select_related('veiculo')

    gastos_despesa = Despesa.objects.filter(
        data__month=mes,
        data__year=ano
    ).select_related('veiculo')

    total_receitas = receitas.aggregate(t=Sum('valor'))['t'] or Decimal('0')
    total_manutencao = gastos_manutencao.aggregate(t=Sum('valor'))['t'] or Decimal('0')
    total_despesa = gastos_despesa.aggregate(t=Sum('valor'))['t'] or Decimal('0')
    total_gastos = total_manutencao + total_despesa
    lucro = total_receitas - total_gastos

    ctx = {
        'receitas': receitas,
        'gastos_manutencao': gastos_manutencao,
        'gastos_despesa': gastos_despesa,
        'total_receitas': total_receitas,
        'total_gastos': total_gastos,
        'lucro': lucro,
        'mes': mes,
        'ano': ano,
    }
    return render(request, 'gestao/financeiro.html', ctx)


@login_required
def relatorios(request):
    opcoes = [
        {'tipo': 'clientes', 'label': 'Clientes', 'icon': '👥', 'desc': 'Lista completa de clientes cadastrados'},
        {'tipo': 'veiculos', 'label': 'Veículos', 'icon': '🚗', 'desc': 'Frota completa com status atual'},
        {'tipo': 'financeiro', 'label': 'Financeiro', 'icon': '💰', 'desc': 'Extrato de receitas e pagamentos'},
        {'tipo': 'manutencoes', 'label': 'Manutenções', 'icon': '🔧', 'desc': 'Histórico de manutenções e gastos'},
        {'tipo': 'pagamentos', 'label': 'Pagamentos', 'icon': '📋', 'desc': 'Status de todos os pagamentos'},
    ]
    return render(request, 'gestao/relatorios.html', {'relatorios_opcoes': opcoes})


@login_required
def relatorio_pdf(request, tipo):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    import io

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []

    hoje = date.today()
    elements.append(Paragraph(f'Relatório de {tipo.title()} - {hoje.strftime("%d/%m/%Y")}', styles['Title']))
    elements.append(Spacer(1, 12))

    if tipo == 'clientes':
        data = [['Nome', 'Sobrenome', 'CPF', 'Endereço']]
        for c in Cliente.objects.filter(ativo=True):
            data.append([c.nome, c.sobrenome, c.cpf, c.endereco])

    elif tipo == 'veiculos':
        data = [['Placa', 'Nome/Modelo', 'Cor', 'Status']]
        for v in Veiculo.objects.filter(ativo=True):
            data.append([v.placa, f'{v.nome} {v.modelo}', v.cor, v.get_status_display()])

    elif tipo == 'pagamentos':
        data = [['Cliente', 'Valor', 'Vencimento', 'Situação']]
        for p in Pagamento.objects.select_related('locacao__cliente').all()[:100]:
            data.append([p.locacao.cliente.nome, f'R$ {p.valor}', str(p.data_vencimento), p.get_situacao_display()])

    elif tipo == 'manutencoes':
        data = [['Veículo', 'Tipo', 'Data', 'Valor']]
        for m in Manutencao.objects.select_related('veiculo').all()[:100]:
            data.append([str(m.veiculo), m.get_tipo_display(), str(m.data), f'R$ {m.valor}'])

    elif tipo == 'financeiro':
        data = [['Cliente', 'Veículo', 'Valor', 'Vencimento', 'Situação']]
        for p in Pagamento.objects.select_related('locacao__cliente', 'locacao__veiculo').all()[:100]:
            data.append([
                p.locacao.cliente.nome,
                str(p.locacao.veiculo),
                f'R$ {p.valor}',
                str(p.data_vencimento),
                p.get_situacao_display(),
            ])

    else:
        data = [['Sem dados']]

    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F2744')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F0F4F8')]),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(table)
    doc.build(elements)

    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="relatorio_{tipo}.pdf"'
    return response


@login_required
def relatorio_excel(request, tipo):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    import io

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = tipo.title()

    header_fill = PatternFill(start_color='0F2744', end_color='0F2744', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)

    if tipo == 'clientes':
        headers = ['Nome', 'Sobrenome', 'CPF', 'Endereço']
        ws.append(headers)
        for c in Cliente.objects.filter(ativo=True):
            ws.append([c.nome, c.sobrenome, c.cpf, c.endereco])

    elif tipo == 'veiculos':
        headers = ['Placa', 'Nome', 'Modelo', 'Cor', 'Status']
        ws.append(headers)
        for v in Veiculo.objects.filter(ativo=True):
            ws.append([v.placa, v.nome, v.modelo, v.cor, v.get_status_display()])

    elif tipo == 'financeiro':
        headers = ['Cliente', 'Veículo', 'Valor', 'Vencimento', 'Pagamento', 'Situação']
        ws.append(headers)
        for p in Pagamento.objects.select_related('locacao__cliente', 'locacao__veiculo').all():
            ws.append([
                p.locacao.cliente.nome,
                str(p.locacao.veiculo),
                float(p.valor),
                str(p.data_vencimento),
                str(p.data_pagamento or ''),
                p.get_situacao_display(),
            ])

    elif tipo == 'pagamentos':
        headers = ['Cliente', 'Veículo', 'Valor', 'Vencimento', 'Pagamento', 'Situação']
        ws.append(headers)
        for p in Pagamento.objects.select_related('locacao__cliente', 'locacao__veiculo').all():
            ws.append([
                p.locacao.cliente.nome,
                str(p.locacao.veiculo),
                float(p.valor),
                str(p.data_vencimento),
                str(p.data_pagamento or ''),
                p.get_situacao_display(),
            ])

    elif tipo == 'manutencoes':
        headers = ['Veículo', 'Tipo', 'Descrição', 'Data', 'Valor', 'Oficina']
        ws.append(headers)
        for m in Manutencao.objects.select_related('veiculo').all():
            ws.append([str(m.veiculo), m.get_tipo_display(), m.descricao, str(m.data), float(m.valor), m.oficina])

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="relatorio_{tipo}.xlsx"'
    return response


@login_required
def pecas_lista(request):
    q = request.GET.get('q', '')
    pecas = Peca.objects.all()
    if q:
        pecas = pecas.filter(nome__icontains=q)
    return render(request, 'gestao/pecas_lista.html', {'pecas': pecas, 'q': q})


@login_required
def pecas_form(request, pk=None):
    peca = get_object_or_404(Peca, pk=pk) if pk else None
    form = PecaForm(request.POST or None, instance=peca)
    if form.is_valid():
        form.save()
        messages.success(request, 'Peça salva com sucesso!')
        return redirect('pecas_lista')
    return render(request, 'gestao/pecas_form.html', {'form': form, 'peca': peca})


@login_required
def peca_excluir(request, pk):
    peca = get_object_or_404(Peca, pk=pk)
    if request.method == 'POST':
        peca.delete()
        messages.success(request, 'Peça removida.')
    return redirect('pecas_lista')


def manifest_json(request):
    manifest = {
        "name": "LM Locadora",
        "short_name": "LM Locadora",
        "description": "Sistema de gestão para locadora de veículos",
        "start_url": "/",
        "display": "standalone",
        "background_color": "#0F2744",
        "theme_color": "#0F2744",
        "orientation": "any",
        "icons": [
            {"src": "/static/icons/icon-192.png", "sizes": "192x192", "type": "image/png"},
            {"src": "/static/icons/icon-512.png", "sizes": "512x512", "type": "image/png"},
        ]
    }
    return JsonResponse(manifest, content_type="application/manifest+json")

def service_worker(request):
    sw_content = """
self.addEventListener('install', (e) => {
  self.skipWaiting();
});
self.addEventListener('activate', (e) => {
  self.clients.claim();
});
self.addEventListener('fetch', (e) => {
  e.respondWith(fetch(e.request));
});
"""
    return HttpResponse(sw_content, content_type='application/javascript')